"""Turns an uploaded file (PDF / DOCX / TXT) into plain text.

Resumes arrive in many formats, but everything downstream (LLM parsing and
evidence search) only needs text, so this is the single place that knows about
file formats."""
import io


def extract_text(filename: str, data: bytes) -> str:
    ext = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if ext == "pdf":
        from pypdf import PdfReader
        text = "\n".join((p.extract_text() or "") for p in PdfReader(io.BytesIO(data)).pages)
    elif ext == "docx":
        from docx import Document
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        text = "\n".join(parts)
    elif ext in ("txt", "md", ""):
        text = data.decode("utf-8", errors="replace")
    else:
        raise ValueError(f"Unsupported file type: .{ext} (use PDF, DOCX or TXT)")
    # PostgreSQL text columns reject NUL bytes, so strip them.
    return text.replace("\x00", "").strip()
