"""Builds the .xlsx HR receives. Runs ONLY after human approval.

Sheets: Summary (JD + counts) | Shortlist | All Candidates | Evidence.
The Evidence sheet is the audit trail: one row per candidate x requirement."""
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..config import settings

HEAD = PatternFill("solid", fgColor="1C2430")
YES = PatternFill("solid", fgColor="D6EFE0")
NO = PatternFill("solid", fgColor="F6D5D1")


def _header(ws, cols):
    ws.append(cols)
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), HEAD
        c.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"


def _fit(ws, cap=60):
    for col in ws.columns:
        width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(width + 2, cap)


def _people_sheet(ws, results, req_names):
    _header(ws, ["Rank", "Name", "Email", "Phone", "Score", "Experience (yrs)", *req_names, "Decision"])
    for r in results:
        by_req = {c["requirement"]: c for c in r.checks}
        row = [r.rank, r.candidate.name or r.candidate.filename, r.candidate.email, r.candidate.phone,
               r.score, r.experience_years]
        row += [by_req[n]["candidate_value"] if n in by_req else "" for n in req_names]
        row.append(("Shortlisted" if r.shortlisted else "Not shortlisted") + (" (HR override)" if r.overridden else ""))
        ws.append(row)
        for i, n in enumerate(req_names, 7):
            cell = ws.cell(ws.max_row, i)
            if n in by_req:
                cell.fill = YES if by_req[n]["met"] else NO
    _fit(ws)


def build_workbook(screening) -> str:
    results = screening.results
    req = screening.requirements
    names = [*req.get("must_have_skills", [])]
    if req.get("min_experience_years"):
        names.append("Experience")
    names += req.get("nice_to_have_skills", [])

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    short = [r for r in results if r.shortlisted]
    for row in [
        ("Role", req.get("role_title") or screening.title), ("HR request", screening.hr_request),
        ("Must-have skills", ", ".join(req.get("must_have_skills", []))),
        ("Nice-to-have skills", ", ".join(req.get("nice_to_have_skills", []))),
        ("Minimum experience", f"{req['min_experience_years']:g}+ years" if req.get("min_experience_years") else "Not specified"),
        ("Candidates analysed", len(results)), ("Shortlisted (after HR review)", len(short)),
        ("Generated", datetime.now().strftime("%Y-%m-%d %H:%M")),
    ]:
        ws.append(row)
        ws.cell(ws.max_row, 1).font = Font(bold=True)
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 30, 90
    for r in ws.iter_rows():
        r[1].alignment = Alignment(wrap_text=True, vertical="top")

    _people_sheet(wb.create_sheet("Shortlist"), short, names)
    _people_sheet(wb.create_sheet("All Candidates"), results, names)

    ev = wb.create_sheet("Evidence")
    _header(ev, ["Candidate", "Requirement", "Required", "Required value", "Candidate", "Evidence", "Evidence type"])
    for r in results:
        for c in r.checks:
            ev.append([r.candidate.name or r.candidate.filename, c["requirement"],
                       "Yes" if c["required"] else "No", c.get("required_value", "Yes" if c["required"] else "Preferred"),
                       c["candidate_value"], c["evidence"] or "— not found in resume —", c.get("evidence_type") or ""])
            ev.cell(ev.max_row, 5).fill = YES if c["met"] else NO
    _fit(ev, cap=90)

    out = Path(settings.export_dir)
    out.mkdir(parents=True, exist_ok=True)
    path = out / f"screening_{screening.id}.xlsx"
    wb.save(path)
    return str(path)
